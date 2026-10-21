"""Audit supplied materials and upstream data; never treats author data as new training."""
import argparse
import ast
import csv
import hashlib
import json
import subprocess
import sys
from collections import defaultdict
from pathlib import Path

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import numpy as np
from openpyxl import load_workbook
from rdkit import Chem, RDLogger

RDLogger.DisableLog('rdApp.*')
p = argparse.ArgumentParser()
p.add_argument('--upstream', type=Path, required=True)
p.add_argument('--materials', type=Path, required=True)
p.add_argument('--out', type=Path, required=True)
a = p.parse_args()
a.out.mkdir(parents=True, exist_ok=True)

def sha(path):
    h = hashlib.sha256()
    with path.open('rb') as f:
        for b in iter(lambda: f.read(1048576), b''):
            h.update(b)
    return h.hexdigest()

materials = [{'name': f.name, 'bytes': f.stat().st_size, 'sha256': sha(f)}
             for f in sorted(a.materials.iterdir()) if f.is_file()]
upstream_sha = subprocess.check_output(['git', '-C', str(a.upstream), 'rev-parse', 'HEAD'], text=True).strip()
counts = []
for f in sorted((a.upstream / 'MotiL_micromolecule/data').glob('*.csv')):
    with f.open(newline='') as stream:
        reader = csv.reader(stream)
        header = next(reader)
        rows = list(reader)
    item = {'dataset': f.stem, 'rows': len(rows), 'label_columns': len(header)-1,
            'header': header, 'sha256': sha(f)}
    if f.stem == 'bbbp':
        mols = [Chem.MolFromSmiles(r[0]) for r in rows]
        valid = [Chem.MolToSmiles(m) for m in mols if m is not None and m.GetNumHeavyAtoms() > 0]
        item.update(valid_rdkit=len(valid), invalid_or_empty=len(rows)-len(valid),
                    unique_canonical=len(set(valid)))
        canonical_rows = defaultdict(list)
        for row_index, (row, mol) in enumerate(zip(rows, mols)):
            if mol is not None and mol.GetNumHeavyAtoms() > 0:
                canonical_rows[Chem.MolToSmiles(mol)].append((row_index, row[1]))
        duplicates = [(smiles, group) for smiles, group in canonical_rows.items() if len(group) > 1]
        item['duplicate_canonical_groups'] = len(duplicates)
        item['conflicting_label_canonical_groups'] = sum(len({label for _, label in group}) > 1
                                                       for _, group in duplicates)
        with (a.out / 'bbbp_duplicate_groups.csv').open('w', newline='') as f:
            writer = csv.writer(f)
            writer.writerow(['canonical_smiles', 'source_row_indices_zero_based', 'labels', 'conflicting_labels'])
            for smiles, group in duplicates:
                writer.writerow([smiles, ';'.join(str(i) for i, _ in group),
                                 ';'.join(label for _, label in group), len({label for _, label in group}) > 1])
    counts.append(item)
with (a.out / 'dataset_counts.csv').open('w', newline='') as f:
    writer = csv.DictWriter(f, fieldnames=['dataset', 'rows', 'label_columns', 'sha256'])
    writer.writeheader()
    writer.writerows({k: d[k] for k in writer.fieldnames} for d in counts)

syntax_errors = []
for f in sorted(a.upstream.rglob('*.py')):
    try:
        ast.parse(f.read_text())
    except (SyntaxError, UnicodeError) as e:
        syntax_errors.append({'path': str(f.relative_to(a.upstream)), 'error': str(e)})

book = load_workbook(a.materials / '41467_2025_66685_MOESM4_ESM.xlsx', data_only=True, read_only=True)
sheet_metadata = []
for sheet in book:
    sheet_metadata.append({'name': sheet.title, 'rows': sheet.max_row, 'columns': sheet.max_column})
    with (a.out / (sheet.title.replace(' ', '_') + '.csv')).open('w', newline='') as f:
        csv.writer(f).writerows(sheet.values)

sheet = book['Figure 2a']
datasets = [sheet.cell(r, 2).value for r in range(6, 10)]
methods = [sheet.cell(5, c).value for c in range(3, 7)]
db = np.array([[sheet.cell(r, c).value for c in range(3, 7)] for r in range(6, 10)])
fig, ax = plt.subplots(figsize=(8, 4))
for i, name in enumerate(methods):
    ax.bar(np.arange(4) + (i-1.5)*0.2, db[:, i], width=0.2, label=name)
ax.set_xticks(np.arange(4), datasets)
ax.set_ylabel('Davies-Bouldin index (lower is better)')
ax.set_title('Author Source Data, Figure 2a (replot only)')
ax.legend(fontsize=8)
fig.tight_layout()
fig.savefig(a.out / 'source_figure2a.png', dpi=180)
plt.close(fig)

sheet = book['Figure 2c']
motifs = [sheet.cell(5, c).value for c in range(3, 13)]
sim = np.array([[sheet.cell(r, c).value for c in range(3, 13)] for r in range(6, 16)])
fig, ax = plt.subplots(figsize=(8, 7))
im = ax.imshow(sim, vmin=0, vmax=1, cmap='viridis')
ax.set_xticks(range(10), motifs, rotation=70, ha='right', fontsize=8)
ax.set_yticks(range(10), motifs, fontsize=8)
ax.set_title('Author Source Data, Figure 2c (replot only)')
fig.colorbar(im, ax=ax, label='Representation similarity')
fig.tight_layout()
fig.savefig(a.out / 'source_figure2c.png', dpi=180)
plt.close(fig)

ckpt = a.upstream / 'MotiL_micromolecule/dumped/pre-train/1-model/original_CMPN_0707_0800_12000th_epoch.pkl'
import torch
summary = {
    'upstream_commit': upstream_sha, 'python': sys.version.split()[0],
    'torch': torch.__version__, 'cuda_available': torch.cuda.is_available(),
    'rdkit': __import__('rdkit').__version__,
    'materials': materials, 'datasets': counts, 'source_data_sheets': sheet_metadata,
    'syntax_errors': syntax_errors,
    'checkpoint': {'relative_path': str(ckpt.relative_to(a.upstream)), 'bytes': ckpt.stat().st_size, 'sha256': sha(ckpt)},
    'figure2a_db_source_data': {'datasets': datasets, 'methods': methods, 'values': db.tolist()},
    'discrepancies': [
        'Main text p.5: QM8 no pretraining DB=5.42; Source Data Figure 2a ESM C9: 6.9. Source Data BBBP C6 is 5.43.',
        'Main Table 2 QM7 count=7160; Supplementary Table 1 QM7 count=6830; inspect dataset_counts.csv for current CSV.',
        'Main Table 1 MUV count=93807; Supplementary Table 1 count=93087; inspect current CSV count.',
        'Main Table 1 ToxCast count=8575; Supplementary Table 1 count=8597; inspect current CSV count.',
        'Source Data Figure 4b cell A2 says Figure 4a: retained verbatim in CSV.',
    ],
    'interpretation': 'Source Data exports and plots are author-result reanalysis, not a new embedding or training replication.',
}
(a.out / 'audit_summary.json').write_text(json.dumps(summary, ensure_ascii=False, indent=2)+'\n')
print(json.dumps({k: summary[k] for k in ['upstream_commit', 'python', 'torch', 'cuda_available', 'syntax_errors']}, indent=2))
print('Dataset counts:', [(x['dataset'], x['rows'], x['label_columns']) for x in counts])
print('Materials:', len(materials), 'Source Data worksheets:', len(sheet_metadata))
